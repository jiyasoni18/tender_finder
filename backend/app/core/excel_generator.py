import json
import re
import logging
from pathlib import Path
from app.core.models import TenderDoc
from app.core.config import REPORTS_DIR

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment
except ImportError:
    openpyxl = None

log = logging.getLogger("excel_generator")

def generate_boq_excel(doc: TenderDoc) -> None:
    if not openpyxl:
        log.warning("openpyxl not installed. Cannot generate BOQ Excel.")
        return
        
    if not doc.report_html:
        return
    
    # Extract JSON from script tag
    match = re.search(r'<script id="boq-json" type="application/json">\s*(.*?)\s*</script>', doc.report_html, re.DOTALL | re.IGNORECASE)
    if not match:
        log.info("No BOQ JSON found in report for %s", doc.doc_id)
        return
        
    try:
        boq_data = json.loads(match.group(1).strip())
    except Exception as e:
        log.error("Failed to parse BOQ JSON for %s: %s", doc.doc_id, e)
        return
        
    if not boq_data:
        log.info("Empty BOQ JSON for %s", doc.doc_id)
        return
        
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOQ"
    
    # Headers
    bold_font = Font(bold=True)
    
    ws.append([f"Tender No: {doc.doc_id}"])
    ws.cell(row=1, column=1).font = bold_font
    
    headers = [
        "S.No.", "Item Code", "Item Qty", "Qty Unit", "Unit Rate", "Basic Value", 
        "Escl.(%)", "Amount", "NIT AMOUNT", "EXPECTED VENDORS RATE", 
        "INCLUDING GST RATE", "VENDOR'S AMOUNT", "PROFIT/LOSS"
    ]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        ws.cell(row=2, column=col_idx).font = bold_font
        
    for item in boq_data:
        # Row 1 for item
        s_no = item.get("s_no", "")
        qty = item.get("qty", 0)
        unit = item.get("unit", "")
        rate = item.get("rate", 0)
        basic_value = item.get("basic_value", 0)
        
        ws.append([
            s_no, None, qty, unit, rate, basic_value, None, basic_value, basic_value, None, None, None, None
        ])
        
        # Row 2 for item: Description
        desc = item.get("description", "")
        ws.append([f"Description:- {desc}"])
        
    # Set column widths
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 10
    ws.column_dimensions['H'].width = 12
    ws.column_dimensions['I'].width = 12
    ws.column_dimensions['J'].width = 22
    ws.column_dimensions['K'].width = 20
    ws.column_dimensions['L'].width = 20
    ws.column_dimensions['M'].width = 15
    
    if doc.pdf_path and doc.pdf_path.parent.is_dir():
        out_dir = doc.pdf_path.parent
    else:
        out_dir = REPORTS_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    
    safe_id = "".join(c if c.isalnum() or c in "-_" else "_" for c in str(doc.doc_id))[:80]
    excel_path = out_dir / f"BOQ_{safe_id}.xlsx"
    
    try:
        wb.save(str(excel_path))
        log.info("✓ BOQ Excel saved to %s", excel_path)
    except Exception as e:
        log.error("Failed to save BOQ Excel for %s: %s", doc.doc_id, e)
