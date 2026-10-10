import openpyxl

def find_columns(sheet):
    """Attempt to find Document Name, Requirement, and Status columns."""
    doc_col = None
    req_col = None
    status_col = None
    notes_col = None
    upload_num_col = None
    
    header_row = 1
    
    # Step 1: Find the first row with actual data (skip headers and blank rows)
    data_start_row = 2
    for row in range(2, min(10, sheet.max_row + 1)):
        has_data = False
        for col in range(1, min(sheet.max_column + 1, 10)):
            val = sheet.cell(row=row, column=col).value
            if val is not None and str(val).strip():
                has_data = True
                break
        if has_data:
            data_start_row = row
            break
    
    # Step 2: Determine doc_col by finding the column with the longest average text
    # in the first few data rows — that's the document description column
    col_avg_len = {}
    sample_end = min(data_start_row + 5, sheet.max_row + 1)
    for col in range(1, min(sheet.max_column + 1, 10)):
        lengths = []
        for row in range(data_start_row, sample_end):
            val = sheet.cell(row=row, column=col).value
            if val is not None:
                lengths.append(len(str(val).strip()))
        if lengths:
            col_avg_len[col] = sum(lengths) / len(lengths)
        else:
            col_avg_len[col] = 0
    
    # The document column has the longest text
    if col_avg_len:
        doc_col = max(col_avg_len, key=col_avg_len.get)
    else:
        doc_col = 2  # fallback
    
    # Step 3: Look for a separate "Required" / "Y/N" / "Mandatory" column
    # Search header rows for keywords
    for row in range(1, min(5, sheet.max_row + 1)):
        for col in range(1, min(sheet.max_column + 1, 10)):
            val = str(sheet.cell(row=row, column=col).value or "").lower().strip()
            if any(kw in val for kw in ["required", "mandatory", "y/n"]):
                if col != doc_col:
                    req_col = col
    
    # If no separate req column found, set it equal to doc_col
    # (means all rows are treated as mandatory)
    if not req_col:
        req_col = doc_col
    
    # Step 4: Search for existing status/notes/upload_num columns in header row
    for col in range(1, sheet.max_column + 1):
        val = str(sheet.cell(row=header_row, column=col).value or "").lower().strip()
        if "document found" in val or "upload status" in val or "found (yes/no)" in val or "found?" in val:
            status_col = col
        elif "notes" in val or "missing details" in val or "description & notes" in val:
            notes_col = col
        elif "upload number" in val or "generated file" in val or "file name" in val:
            upload_num_col = col
            
    # Step 5: Determine next available column for status output
    # Find the last column with actual data
    last_data_col = max(doc_col, req_col)
    for col in range(1, sheet.max_column + 1):
        for row in range(data_start_row, min(data_start_row + 3, sheet.max_row + 1)):
            if sheet.cell(row=row, column=col).value is not None:
                last_data_col = max(last_data_col, col)
                break
    
    next_empty_col = last_data_col + 1
    
    # Step 6: Unmerge any merged cells that might block our status columns
    for mc in list(sheet.merged_cells.ranges):
        try:
            sheet.unmerge_cells(str(mc))
        except:
            pass
    
    # Step 7: Create status columns if they don't exist
    if not status_col:
        status_col = next_empty_col
        sheet.cell(row=header_row, column=status_col).value = "Document Found?"
        next_empty_col += 1
    if not upload_num_col:
        upload_num_col = next_empty_col
        sheet.cell(row=header_row, column=upload_num_col).value = "Generated File Name"
        next_empty_col += 1
    if not notes_col:
        notes_col = next_empty_col
        sheet.cell(row=header_row, column=notes_col).value = "Description & Notes"
        
    return header_row, doc_col, req_col, status_col, upload_num_col, notes_col

def parse_excel_requirements(file_path):
    print(f"Reading Excel: {file_path}")
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    
    header_row, doc_col, req_col, status_col, upload_num_col, notes_col = find_columns(sheet)
    
    requirements = []
    
    # Find the actual data start row (skip empty rows after header)
    data_start = header_row + 1
    for row in range(header_row + 1, min(header_row + 5, sheet.max_row + 1)):
        val = sheet.cell(row=row, column=doc_col).value
        if val and str(val).strip():
            data_start = row
            break
    
    for row in range(data_start, sheet.max_row + 1):
        doc_name = sheet.cell(row=row, column=doc_col).value
        
        if doc_name and str(doc_name).strip():
            is_optional = False
            # Only check the requirement column if it's a separate column from the document name
            if req_col != doc_col:
                req_val = str(sheet.cell(row=row, column=req_col).value or "").lower().strip()
                is_optional = "no" in req_val.split() or "optional" in req_val or "not required" in req_val
            
            if not is_optional:
                requirements.append({
                    "row": row,
                    "name": str(doc_name).strip()
                })
                
    wb.close()
    print(f"  Found {len(requirements)} mandatory requirements (doc_col={doc_col}, req_col={req_col}, status_col={status_col})")
    return requirements, header_row, doc_col, req_col, status_col, upload_num_col, notes_col

def update_excel_status(file_path, results, status_col, upload_num_col, notes_col):
    """
    results is a dict of row_index -> {"found": True/False, "notes": "...", "upload_num": "01_..."}
    """
    print(f"Updating Excel with results: {file_path}")
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    
    for row, data in results.items():
        if not isinstance(row, int):
            continue
        found = data.get("found", False)
        notes = data.get("notes", "")
        upload_num = data.get("upload_num", "")
        
        sheet.cell(row=row, column=status_col).value = "YES" if found else "NO"
        sheet.cell(row=row, column=upload_num_col).value = upload_num if upload_num else ""
        sheet.cell(row=row, column=notes_col).value = notes if notes else ""
            
    try:
        wb.save(file_path)
        print(f"  Excel saved successfully.")
    except PermissionError:
        print(f"\n[!] ERROR: Could not save to {file_path}")
        print("Please close the Excel file if it is open in Microsoft Excel and try again.")
    wb.close()

def replace_excel_requirements(file_path, new_reqs, header_row, doc_col):
    """
    Clears existing requirements in the excel file and inserts the new_reqs (from PDF).
    Returns the formatted requirements list with row numbers.
    """
    print(f"Replacing Excel requirements with NIT requirements in: {file_path}")
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    
    if sheet.max_row > header_row:
        sheet.delete_rows(header_row + 1, sheet.max_row - header_row)
        
    reqs_format = []
    for idx, req_text in enumerate(new_reqs):
        row_num = header_row + 1 + idx
        sheet.cell(row=row_num, column=doc_col).value = req_text
        reqs_format.append({
            "row": row_num,
            "name": req_text
        })
        
    try:
        wb.save(file_path)
    except Exception as e:
        print(f"Error saving updated excel: {e}")
    wb.close()
    return reqs_format

def create_excel_from_requirements(file_path, reqs):
    """
    Creates a new excel file with the given requirements.
    Returns excel_reqs, header_row, doc_col, req_col, status_col, upload_num_col, notes_col
    """
    print(f"Creating new Excel mapping file: {file_path}")
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Requirements"
    
    headers = ["S.No", "Requirement", "Status", "Upload Number", "Notes"]
    for col, h in enumerate(headers, 1):
        sheet.cell(row=1, column=col).value = h
        
    reqs_format = []
    for idx, req_text in enumerate(reqs):
        row_num = idx + 2
        sheet.cell(row=row_num, column=1).value = idx + 1
        sheet.cell(row=row_num, column=2).value = req_text
        reqs_format.append({"row": row_num, "name": req_text})
        
    wb.save(file_path)
    wb.close()
    
    return reqs_format, 1, 2, 2, 3, 4, 5
