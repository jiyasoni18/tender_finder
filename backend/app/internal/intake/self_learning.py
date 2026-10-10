import os
import yaml
import re
import openpyxl
from pathlib import Path

try:
    from google import genai
    HAS_GEMINI = True
except ImportError:
    HAS_GEMINI = False

BASE_DIR = Path(__file__).resolve().parent.parent.parent.parent
CATALOG_PATH = BASE_DIR / "tender-system" / "app" / "catalog" / "requirements_ireps.yaml"

def load_catalog():
    if not CATALOG_PATH.exists():
        return []
    with open(CATALOG_PATH, 'r') as f:
        data = yaml.safe_load(f)
    reqs = data.get("requirements") if data else None
    return reqs if reqs is not None else []

def save_catalog(catalog):
    with open(CATALOG_PATH, 'w') as f:
        yaml.dump({"requirements": catalog}, f, default_flow_style=False, sort_keys=False)

def extract_canonical_ids_from_filename(filename, catalog):
    if not filename:
        return []
        
    import re
    # Split by common delimiters the user might use
    parts = re.split(r'[,+&|/]', filename)
    
    matched_ids = set()
    
    for part in parts:
        # Clean filename by removing prefix like "01_" or "12_"
        cleaned_name = re.sub(r'^\d+_', '', part).strip().lower()
        if not cleaned_name:
            continue
            
        for req in catalog:
            if req.get("filename_hint") and req["filename_hint"].lower() in cleaned_name:
                matched_ids.add(req["id"])
            
            if req["id"].lower() in cleaned_name:
                matched_ids.add(req["id"])
                
            for alias in req.get("aliases", []):
                alias_lower = alias.lower()
                # Exact match or word boundary match for short aliases (like aoa, moa, jv, poa)
                if len(alias_lower) <= 4:
                    if re.search(r'\b' + re.escape(alias_lower) + r'\b', cleaned_name):
                        matched_ids.add(req["id"])
                else:
                    # Substring match for longer aliases
                    if alias_lower in cleaned_name or cleaned_name in alias_lower:
                        matched_ids.add(req["id"])
                        
    return list(matched_ids)

def find_unmapped_corrections_in_excel(excel_path, catalog, map_with_keywords_func):
    try:
        wb = openpyxl.load_workbook(excel_path)
    except Exception as e:
        print(f"Error reading {excel_path}: {e}")
        return []
        
    sheet = wb.active
    
    # Need to find the columns
    from app.internal.intake.excel_parser import find_columns
    header_row, doc_col, req_col, status_col, upload_num_col, notes_col = find_columns(sheet)
    
    unmapped = []
    
    for row in range(header_row + 1, sheet.max_row + 1):
        requirement = sheet.cell(row=row, column=doc_col).value
        generated_file = sheet.cell(row=row, column=upload_num_col).value
        
        if not requirement or not generated_file:
            continue
            
        requirement = str(requirement).strip()
        generated_file = str(generated_file).strip()
        
        canonical_ids = extract_canonical_ids_from_filename(generated_file, catalog)
        
        if canonical_ids:
            current_mapped_ids = map_with_keywords_func(requirement, catalog)
            
            for cid in canonical_ids:
                if cid not in current_mapped_ids:
                    unmapped.append({
                        "requirement": requirement,
                        "canonical_id": cid,
                        "excel_path": str(excel_path)
                    })
    wb.close()
    return unmapped

def batch_learn_keywords(unmapped_list, api_key):
    """
    Takes a list of dicts: {"requirement": ..., "canonical_id": ...}
    Returns a list of dicts with {"canonical_id": ..., "new_alias": ...}
    Uses a single Gemini call.
    """
    if not unmapped_list or not api_key or not HAS_GEMINI:
        return []
        
    client = genai.Client(api_key=api_key)
    
    prompt = "You are an AI assistant helping to improve a document mapping system. "
    prompt += "I will give you a list of requirement texts and the ID of the document they should map to. "
    prompt += "For each item, extract a short, concise noun phrase (1-4 words) from the requirement text that acts as the core keyword/alias for this document. "
    prompt += "Do not invent words; use exact words/phrases from the requirement text if possible. "
    prompt += "Format your response exactly as a JSON array of objects, e.g.:\n"
    prompt += '[{"index": 0, "alias": "financial turnover"}, {"index": 1, "alias": "experience certificate"}]\n\n'
    
    prompt += "Items to process:\n"
    for idx, item in enumerate(unmapped_list):
        prompt += f"Index: {idx}\nRequirement: {item['requirement']}\nCanonical ID: {item['canonical_id']}\n\n"
        
    try:
        response = client.models.generate_content(
            model='gemini-3.1-flash-lite',
            contents=prompt,
        )
        
        response_text = response.text.strip()
        # Clean up JSON if it has markdown ticks
        if response_text.startswith("```json"):
            response_text = response_text[7:]
        if response_text.endswith("```"):
            response_text = response_text[:-3]
            
        import json
        learned_data = json.loads(response_text)
        
        results = []
        for data in learned_data:
            idx = data.get("index")
            alias = data.get("alias")
            if idx is not None and alias and 0 <= idx < len(unmapped_list):
                results.append({
                    "canonical_id": unmapped_list[idx]["canonical_id"],
                    "new_alias": alias.lower().strip()
                })
        return results
    except Exception as e:
        print(f"Error during Gemini batch learning: {e}")
        return []

def apply_new_aliases(new_aliases, catalog):
    updated = False
    for item in new_aliases:
        cid = item["canonical_id"]
        alias = item["new_alias"]
        
        for req in catalog:
            if req["id"] == cid:
                if "aliases" not in req:
                    req["aliases"] = []
                if alias not in req["aliases"]:
                    req["aliases"].append(alias)
                    updated = True
                    print(f"[*] Learned new alias for {cid}: '{alias}'")
                break
    return updated

def run_learning_cycle(tenders_dir=None):
    from app.internal.intake.intake import map_with_keywords
    
    if tenders_dir is None:
        tenders_dir = BASE_DIR / "data" / "tender-system" / "tenders"
        
    if not tenders_dir.exists():
        print("Tenders directory not found.")
        return
        
    catalog = load_catalog()
    if not catalog:
        print("Catalog not found.")
        return
        
    print(f"Scanning for corrections in {tenders_dir}...")
    
    all_unmapped = []
    
    from pathlib import Path
    for excel_path in Path(tenders_dir).rglob("Mapped_Requirements_*.xlsx"):
        unmapped = find_unmapped_corrections_in_excel(excel_path, catalog, map_with_keywords)
        if unmapped:
            all_unmapped.extend(unmapped)
            
    if not all_unmapped:
        print("No user corrections found that need learning.")
        return
        
    print(f"Found {len(all_unmapped)} potential corrections to learn from.")
    
    # Deduplicate unmapped requirements to save tokens
    seen_reqs = set()
    unique_unmapped = []
    for item in all_unmapped:
        if item["requirement"] not in seen_reqs:
            seen_reqs.add(item["requirement"])
            unique_unmapped.append(item)
            
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        print("GEMINI_API_KEY not found. Cannot run learning.")
        return
        
    print(f"Batching {len(unique_unmapped)} items for Gemini extraction...")
    
    # Process in batches of 10 to avoid huge prompts
    batch_size = 10
    learned_aliases = []
    for i in range(0, len(unique_unmapped), batch_size):
        batch = unique_unmapped[i:i+batch_size]
        new_aliases = batch_learn_keywords(batch, api_key)
        learned_aliases.extend(new_aliases)
        
    if learned_aliases:
        updated = apply_new_aliases(learned_aliases, catalog)
        if updated:
            save_catalog(catalog)
            print("Successfully updated requirements_ireps.yaml with new learnings.")
        else:
            print("No new unique aliases were generated.")
    else:
        print("No aliases were extracted.")

def check_and_trigger_learning():
    tenders_dir = BASE_DIR / "data" / "tender-system" / "tenders"
    if not tenders_dir.exists(): return
    
    folders = [d for d in tenders_dir.iterdir() if d.is_dir()]
    num_folders = len(folders)
    
    state_file = BASE_DIR / "data" / "tender-system" / ".learning_state"
    last_learned = 0
    if state_file.exists():
        try:
            last_learned = int(state_file.read_text().strip())
        except:
            pass
            
    if num_folders >= last_learned + 5:
        print(f"Triggering learning cycle (folders={num_folders}, last_learned={last_learned})")
        import threading
        def run_bg():
            run_learning_cycle(tenders_dir)
            try:
                state_file.parent.mkdir(parents=True, exist_ok=True)
                with open(state_file, 'w') as f:
                    f.write(str(num_folders))
            except Exception as e:
                print(f"Failed to update learning state: {e}")
                
        threading.Thread(target=run_bg, daemon=True).start()

