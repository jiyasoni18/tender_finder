import markdown
import asyncio
import logging
import threading
import urllib.parse
import re
import requests
from pathlib import Path

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from fastapi import FastAPI, WebSocket, WebSocketDisconnect, Request
from app.core.pdf_generator import generate_pdf_from_html
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse

from app.core.config import SITES, PIPELINE
import app.core.config as config
from app.core.db import init_db
from app.core.state import Pipeline, Ledger

from google import genai
gemini_client = None
if config.GEMINI_API_KEY:
    gemini_client = genai.Client(api_key=config.GEMINI_API_KEY)
else:
    logging.warning("GEMINI_API_KEY is not set.")
from fastapi import APIRouter, File, UploadFile, Request, Form, Response, Depends, BackgroundTasks

router = APIRouter()

@router.post("/api/summary/batch")
async def generate_batch_summaries():
    """Automatically generate summaries for all passed tenders that lack one."""
    if not config.GEMINI_API_KEY:
        return {"status": "error", "message": "GEMINI_API_KEY not configured."}
        
    def get_md_content(comp: str):
        if comp.lower() == "arkonic":
            path = Path(r"c:\Users\DELL\Downloads\tender_finder\tender_finder\GEM_DOCUMENT_LIBRARY.md")
            return path.read_text(encoding="utf-8") if path.exists() else ""
        return ""
        
    md_content = get_md_content("arkonic") # Batch uses Arkonic by default for now
    
    prompt = f"""
    You are an expert procurement assistant.
    Read the attached Tender Document (NIT) and create a comprehensive summary in Markdown format.
    Crucially, cross-reference the requirements in the tender against our company's Document Library rules provided below.
    
    Document Library Rules:
    {md_content}
    
    Format your response EXACTLY mimicking this strict structure (use markdown tables):
    
    01. TENDER OVERVIEW
    | Field | Details |
    |---|---|
    | Bid Number | ... |
    | Bid Document Dated | ... |
    | Department / Ministry | ... |
    | Item Category & Quantity | ... |
    | Bid End Date | ... |
    | Estimated Bid Value (ECV) | ... |
    | EMD & ePBG | ... |
    
    02. KEY TERMS & SCOPE
    (List any important supply and installation scopes, SLAs, delivery timelines)
    
    03. MANDATORY SUBMISSION CHECKLIST & DOCUMENT MAPPING
    | # | Document as named in the tender | Action & file | Review |
    |---|---|---|---|
    (For Action & file, use 'MAP: [file path]' if the file is in our library, 'CREATE' if we must make it, or 'MANUAL REQUIREMENT' if external. Provide thoughtful reviews for each.)
    """
    
    def bg_summarize():
        ledger = _get_ledger()
        passed = ledger.get_passed_tenders()
        
        for row in passed:
            tender_id = row["id"]
            source = row.get("source", "")
            tender_dir = config.DOWNLOADS_DIR / source / tender_id
            if not tender_dir.exists():
                tender_dir = config.DOWNLOADS_DIR / tender_id
            
            summary_name = f"Summary_{tender_id}.md"
            summary_path = tender_dir / "summary" / summary_name
            if summary_path.exists():
                continue
                
            pdf_path = None
            if tender_dir.is_dir():
                for f in tender_dir.iterdir():
                    if f.suffix.lower() == ".pdf" and not f.name.startswith("Summary_"):
                        pdf_path = f
                        break
            
            if not pdf_path:
                ireps_dir = Path(r"C:\Users\DELL\Downloads\IREPS_Tenders") / tender_id
                if ireps_dir.is_dir():
                    for f in ireps_dir.iterdir():
                        if f.suffix.lower() == ".pdf" and not f.name.startswith("Summary_"):
                            pdf_path = f
                            summary_path = ireps_dir / "summary" / summary_name
                            break
                            
            if pdf_path:
                try:
                    logging.info(f"Generating AI Summary for {tender_id}...")
                    atc_path = None
                    try:
                        doc = fitz.open(str(pdf_path))
                        for page in doc:
                            for link in page.get_links():
                                uri = link.get("uri", "")
                                if uri and ("get_atc_document" in uri.lower() or "atc" in uri.lower() or "slafds" in uri.lower()):
                                    atc_res = requests.get(uri, verify=False, timeout=(5, 10))
                                    if atc_res.status_code == 200:
                                        ext = ".pdf"
                                        if ".docx" in uri.lower(): ext = ".docx"
                                        elif ".doc" in uri.lower(): ext = ".doc"
                                        base_name = pdf_path.name.replace(".pdf", "")
                                        atc_path = pdf_path.with_name(f"ATC_{base_name}{ext}")
                                        atc_path.write_bytes(atc_res.content)
                                    break
                            if atc_path:
                                break
                    except Exception as e:
                        logging.warning(f"Could not extract ATC link for {tender_id}: {e}")
                        
                    uploaded_files = []
                    uploaded_files.append(gemini_client.files.upload(file=str(pdf_path)))
                    if atc_path and atc_path.exists():
                        uploaded_files.append(gemini_client.files.upload(file=str(atc_path)))
                        
                    bg_prompt = f"""
                    You are an expert procurement assistant.
                    Read the attached Tender Document (NIT){" and the attached ATC document" if atc_path else ""} and create a comprehensive summary in Markdown format.
                    Crucially, cross-reference the requirements in the tender against our company's Document Library rules provided below.
                    
                    {"IMPORTANT: I have provided the ATC document. The REAL checklist for the bidder is often located in the ATC document under 'Section-1 (Part B) [Check List for Bidder]'. Make sure you use the ATC checklist for mapping!" if atc_path else ""}
                    
                    Document Library Rules:
                    {get_md_content("arkonic")}
                    
                    Format your response EXACTLY mimicking this strict 6-section structure (use markdown tables):
                    
                    01. TENDER OVERVIEW
                    | Field | Details |
                    |---|---|
                    | Bid Number | ... |
                    | Bid Document Dated | ... |
                    | Department / Ministry | ... |
                    | Item Category & Quantity | ... |
                    | Bid End Date | ... |
                    | Estimated Bid Value (ECV) | ... |
                    
                    02. KEY TERMS & SCOPE
                    (List any important supply and installation scopes, SLAs, delivery timelines)
                    
                    03. ELIGIBILITY CRITERIA
                    (List the exact past experience, OEM authorizations, or general eligibility needed)
                    
                    04. FINANCIAL REQUIREMENTS
                    (List the Turnover, Net Worth, Solvency, EMD, ePBG, and any other financial limits)
                    
                    05. MANDATORY SUBMISSION CHECKLIST & DOCUMENT MAPPING
                    | # | Document as named in the tender | Action & file | Review |
                    |---|---|---|---|
                    (For Action & file, use 'MAP: [file path]' if the file is in our library, 'CREATE' if we must make it, or 'MANUAL REQUIREMENT' if external. Provide thoughtful reviews for each.)
                    
                    06. IMPORTANT DATES & SUBMISSION DETAILS
                    (List bid start/end dates, opening dates, and any physical submission requirements)
                    """
                    
                    contents = uploaded_files + [bg_prompt]
                    import time
                    FALLBACK_MODELS = ["gemini-3.1-flash-lite", "gemini-3.5-flash-lite", "gemini-3.5-flash", "gemini-3.6-flash", "gemini-3.7-flash"]
                    response = None
                    for model_name in FALLBACK_MODELS:
                        try:
                            logging.info(f"Trying model: {model_name}")
                            response = gemini_client.models.generate_content(model=model_name, contents=contents)
                            logging.info(f"Success with model: {model_name}")
                            break
                        except Exception as e:
                            logging.warning(f"Model {model_name} failed: {str(e)[:80]}")
                            if model_name == FALLBACK_MODELS[-1]:
                                raise e
                            time.sleep(2)
                    
                    html_content = response.text
                    html_content = re.sub(r"^```(?:markdown|html)?\s*", "", html_content, flags=re.MULTILINE)
                    html_content = re.sub(r"```\s*$", "", html_content, flags=re.MULTILINE)
                    html_content = markdown.markdown(html_content, extensions=['tables'])
                    
                    template_path = Path("backend/core/pdf_template.html")
                    if template_path.exists():
                        template = template_path.read_text(encoding="utf-8")
                    else:
                        template = "<html><body>{{CONTENT}}</body></html>"
                        
                    full_html = template.replace("{{TENDER_NO}}", tender_id).replace("{{TENDER_TITLE}}", "Tender Summary").replace("{{CONTENT}}", html_content)
                    
                    summary_path_pdf = summary_path.with_suffix(".pdf")
                    summary_path_pdf.parent.mkdir(parents=True, exist_ok=True)
                    generate_pdf_from_html(full_html, str(summary_path_pdf))
                except Exception as e:
                    logging.error(f"Summary generation failed for {tender_id}: {e}")

    import threading
    threading.Thread(target=bg_summarize, daemon=True).start()
    return {"status": "success", "message": "Batch summarization started."}


def get_document_library_and_profile(site: str, comp: str) -> tuple[str, str, str]:
    """
    Returns (company_display_name, short_name, document_library_rules_markdown)
    tailored to the selected portal (IREPS / GeM) and company (Arkonic / Aquint).
    """
    site_lower = site.lower().strip()
    comp_lower = comp.lower().strip()
    root_dir = Path(r"c:\Users\DELL\Downloads\tender_finder\tender_finder")
    
    if "arkon" in comp_lower:
        company_display = "ARKONIK FIBERTECH INDIA PRIVATE LIMITED"
        short_name = "Arkonik"
        gem_lib = root_dir / "GEM_DOCUMENT_LIBRARY.md"
        base_md = gem_lib.read_text(encoding="utf-8") if gem_lib.exists() else ""
        
        if site_lower == "gem":
            md_content = base_md
        else:
            # Arkonic on IREPS / Railway
            md_content = f"""# Railway/IREPS Document Library — ARKONIK FIBERTECH INDIA PRIVATE LIMITED
Entity: ARKONIK FIBERTECH INDIA PRIVATE LIMITED (Ahmedabad, Gujarat)
CIN: U74996GJ2023PTC140363
MSME: Udyam-GJ-01-0352752 (Micro, Services) | DPIIT Startup: DIPP251312
Authorised Signatory / Director: Dharamrajsinh Narendrasinh Vaghela

## Railway/IREPS Mapping Rules for Arkonic:
- Board Resolution / Power of Attorney: MAP: `BR.pdf` (or `POAA FORMATE.odt`)
- Financial Turnover (last 3 years): MAP: `KYC DOCS/ARKONIK TURNOVER DONE.pdf`
- Certificate of Incorporation / CIN: MAP: `KYC DOCS/INC-24_Certificate of Incorporation Pursuant To Change of Name_AA6440272 (1).pdf`
- Similar Work Experience: MAP: `KYC DOCS/arkonik EXPERIENCE.pdf`
- Statutory KYC: PAN (`KYC DOCS/Arkonic PAN.pdf`), GST (`KYC DOCS/ARKONIK GST.pdf`), MSME (`KYC DOCS/Arkonik MSME.pdf`), Startup (`KYC DOCS/ARKONIK STARTUP.pdf`), Balance Sheet (`KYC DOCS/Arkonik Fibetech BSheet 01-04-25 to 31-03-26_CA certified.pdf`), PF/ESIC (`KYC DOCS/Arkonik PF Certificate.pdf`, `KYC DOCS/ESIC certificate......pdf`), Bank Solvency (`Arkonik Solcency SBI.pdf`)
- Relative Declaration: MAP+EDIT: `Near Relation Certificates.docx`
- Statement of Deviations: MAP+EDIT: `statement of deviations.docx` (filled Nil)
- Compliance Certificate: MAP+EDIT: `Certificate from Bidder for compliance.docx`
- Non-Blacklisting / Debarment: MAP: `DECLARATION REGARDING NON BLACKLIST by GST authorities.pdf`
- Other per-tender Railway annexures: CREATE on company letterhead

{base_md}
"""
    else:
        # Aquint
        company_display = "AQUINT INFRATELE INDIA PRIVATE LIMITED"
        short_name = "Aquint"
        rail_lib = root_dir / "aquint" / "RAILWAY_DOCUMENT_LIBRARY.md"
        if not rail_lib.exists():
            rail_lib = root_dir / "AQUIANT_DOCUMENT_LIBRARY.md"
        base_rail_md = rail_lib.read_text(encoding="utf-8") if rail_lib.exists() else ""
        
        if site_lower == "ireps":
            md_content = base_rail_md
        else:
            # Aquint on GeM
            md_content = f"""# GeM Document Library — AQUINT INFRATELE INDIA PRIVATE LIMITED
Entity: AQUINT INFRATELE INDIA PRIVATE LIMITED (Ahmedabad, Gujarat)
MSME: UDYAM-GJ-01-0133628 (General category) | DPIIT Startup Registered
Average Annual Contractual Turnover: ₹22.96 Cr (CA certified Proforma-2)
Telecom / OFC Experience: ₹94.38 Lakhs (GeM/Indian Air Force work order)

## GeM Mapping Rules for Aquint:
- Board Resolution / Power of Attorney: MAP: `11.POWER OF ATTORNEY_BR.pdf` or `POWER OF ATTORNEY.pdf`
- Bidder Turnover: MAP: `TURNOVER (3 YEARS).pdf`
- Certificate of Incorporation / CIN / Constitution: MAP: `5.AOA.MOA.CIN.POA.pdf` or `Certificate_of_Incorporation.pdf`
- Experience / Past Performance: MAP: `EXPERIENCE AQUINT.pdf`, `28.Telecome Work.pdf`
- Statutory KYC: PAN (`PAN_CARD.pdf`), GST (`GST.pdf`), MSME (`MSME.pdf`), Startup (`STARTUPS_CERTIFICATE.pdf`), Audited Balance Sheet (`balance sheet.pdf`), ITR 3 Years (`ITR LAST THREE YEARS.pdf`), EPF (`PF_CERTIFICATE.pdf`), Bank Mandate / Cancelled Cheque (`SBI_CANCEL_CHEQUE.pdf`, `9.MANDATE FORM.docx`)
- Common Undertakings: `2.DECLARATION FORM FOR RELATIVE.docx`, `12.STATEMENT OF DEVIATIONS - FORMAT-1.docx`, `20.Black_listed_by_Govt_Authority.docx`, `18.make in india.docx`, `23.Declaration of Non-Applicability of JV and Partnership.docx`
- Any special ATC formats provided by buyer: MAP directly from buyer ATC format or CREATE on letterhead

{base_rail_md}
"""

    return company_display, short_name, md_content


@router.post("/api/summary/generate")
async def generate_summary(
    file: UploadFile = File(...),
    website: str = Form("ireps"),
    company: str = Form("Aquint"),
    atc_file: UploadFile | None = File(None),
    atc_file_name: str | None = Form(None),
):
    """
    Phase 2: Generate a summary from an uploaded NIT document with tailored website & company document mapping.
    """
    try:
        content = await file.read()
        file_path = config.DOWNLOADS_DIR / file.filename
        file_path.parent.mkdir(parents=True, exist_ok=True)
        file_path.write_bytes(content)
        
        base_name = file.filename.replace(".pdf", "")
        clean_tender_no = re.sub(r'[\\/*?:"<>|]', "_", base_name)
        atc_path = None
        
        if website.lower() == "gem":
            if atc_file:
                atc_content = await atc_file.read()
                atc_path = config.DOWNLOADS_DIR / atc_file.filename
                atc_path.write_bytes(atc_content)
            elif atc_file_name:
                cand = config.DOWNLOADS_DIR / atc_file_name
                if cand.exists():
                    atc_path = cand
                elif (file_path.parent / atc_file_name).exists():
                    atc_path = file_path.parent / atc_file_name
            
            if not atc_path:
                try:
                    from app.routers.docs import download_and_unify_gem_atc
                    unified_atc, count = download_and_unify_gem_atc(file_path, clean_tender_no, orig_stem=base_name)
                    if unified_atc and unified_atc.exists():
                        atc_path = unified_atc
                except Exception as e:
                    logging.warning(f"Could not download unified ATC for summary: {e}")
        else:
            # Fallback link search for IREPS or general tenders
            try:
                doc = fitz.open(str(file_path))
                for page in doc:
                    for link in page.get_links():
                        uri = link.get("uri", "")
                        if uri and ("get_atc_document" in uri.lower() or "atc" in uri.lower() or "slafds" in uri.lower()):
                            headers = { "User-Agent": "Mozilla/5.0", "Accept": "*/*" }
                            atc_res = requests.get(uri, verify=False, timeout=(5, 10), headers=headers)
                            if atc_res.status_code == 200:
                                ext = ".pdf"
                                if ".docx" in uri.lower(): ext = ".docx"
                                elif ".doc" in uri.lower(): ext = ".doc"
                                atc_path = file_path.with_name(f"ATC_{clean_tender_no}{ext}")
                                atc_path.write_bytes(atc_res.content)
                            break
                    if atc_path:
                        break
            except Exception as e:
                logging.warning(f"Could not extract ATC link: {e}")
        
        if not config.GEMINI_API_KEY:
            raise ValueError("GEMINI_API_KEY not configured in backend.")
            
        company_display, short_company, md_content = get_document_library_and_profile(website, company)
        
        # Extract checklist text from ATC document (.docx or .pdf)
        atc_checklist_text = ""
        if atc_path and atc_path.exists():
            atc_str = str(atc_path).lower()
            try:
                if atc_str.endswith(".docx"):
                    # --- DOCX extraction ---
                    from docx import Document as DocxDocument
                    from docx.oxml.ns import qn
                    docx_doc = DocxDocument(str(atc_path))
                    
                    # Find the "Check List for Bidder" table
                    found_checklist = False
                    for elem in docx_doc.element.body:
                        if elem.tag == qn('w:p'):
                            texts = [node.text for node in elem.iter(qn('w:t')) if node.text]
                            text = "".join(texts)
                            if "Check List" in text:
                                found_checklist = True
                        elif elem.tag == qn('w:tbl') and found_checklist:
                            for tbl in docx_doc.tables:
                                if tbl._element is elem:
                                    rows_text = []
                                    for row in tbl.rows:
                                        cells = [cell.text.strip() for cell in row.cells]
                                        rows_text.append(" | ".join(cells))
                                    atc_checklist_text = "\n".join(rows_text)
                                    break
                            break
                    
                    if not atc_checklist_text:
                        # Fallback: extract all text from the docx
                        all_text = []
                        for para in docx_doc.paragraphs:
                            if para.text.strip():
                                all_text.append(para.text.strip())
                        for table in docx_doc.tables:
                            for row in table.rows:
                                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                                if row_text.strip():
                                    all_text.append(row_text)
                        atc_checklist_text = "\n".join(all_text)
                
                elif atc_str.endswith(".pdf"):
                    # --- PDF extraction ---
                    atc_doc = fitz.open(str(atc_path))
                    all_text_parts = []
                    checklist_found = False
                    
                    for page in atc_doc:
                        page_text = page.get_text("text")
                        all_text_parts.append(page_text)
                        
                        # Also extract tables from PDF
                        try:
                            tables = page.find_tables()
                            for table in tables:
                                table_data = table.extract()
                                for row in table_data:
                                    row_text = " | ".join(str(cell).strip() if cell else "" for cell in row)
                                    all_text_parts.append(row_text)
                        except Exception:
                            pass
                    
                    full_atc_text = "\n".join(all_text_parts)
                    
                    # Try to find "Section-10" or "Check List" section
                    found_idx = -1
                    for marker in ["Section-10", "section-10", "Check List", "check list", "Checklist", "checklist", "List of documents to be submitted"]:
                        idx = full_atc_text.find(marker)
                        if idx >= 0:
                            found_idx = idx
                            break

                    if found_idx >= 0:
                        atc_checklist_text = full_atc_text[found_idx:found_idx + 15000]
                    else:
                        atc_checklist_text = full_atc_text[:20000]
                    
                    atc_doc.close()
                
                if atc_checklist_text:
                    logging.info(f"Extracted ATC checklist ({atc_str.split('.')[-1]}): {len(atc_checklist_text)} chars")
                    
            except Exception as e:
                logging.warning(f"Could not extract ATC checklist: {e}")
        
        uploaded_files = []
        uploaded_files.append(gemini_client.files.upload(file=str(file_path)))
        if atc_path and atc_path.exists():
            uploaded_files.append(gemini_client.files.upload(file=str(atc_path)))
        
        atc_section = ""
        if atc_checklist_text:
            atc_section = f"""
        CRITICAL: Below is the EXACT checklist extracted from the ATC document 'Section-10 / Check List for Bidder'.
        You MUST use this EXACT list for Section 06. Do NOT invent or modify document names. Map each item below one-by-one:
        
        --- START OF ATC CHECKLIST ---
        {atc_checklist_text}
        --- END OF ATC CHECKLIST ---
        """
        
        prompt = f"""
        You are an expert procurement consultant producing a professional-grade tender analysis for portal {website.upper()}.
        Read the attached Tender Document (NIT){" and the attached ATC document" if atc_path else ""} thoroughly.
        Cross-reference the tender requirements against our company's Document Library rules provided below.
        
        Our Company: {company_display} ({short_company})
        Bidding Portal: {website.upper()}
        Document Library Rules:
        {md_content}
        
        {atc_section}
        
        Produce a COMPREHENSIVE, DETAILED analysis in Markdown format with EXACTLY these 6 sections.
        Be extremely thorough — extract every single detail from the tender. Do NOT summarize briefly.
        
        ### 01. TENDER OVERVIEW
        Create a detailed table with ALL of the following fields (leave blank only if truly not found in the document):
        | Field | Details |
        |---|---|
        | Bid / Tender Number | ... |
        | Bid Document Dated | ... |
        | Ministry / Department | ... |
        | Organisation / Office | ... |
        | Tender Inviting Authority | Full name, address, designation |
        | Item Category | ... |
        | Similar Category Cluster | ... |
        | Section & Scope Description | Detailed description of what is being procured, locations, quantities |
        | Contract Period | Duration, extension options, termination terms |
        | Bid End Date / Time | ... |
        | Bid Opening Date / Time | ... |
        | Query Deadline | ... |
        | Offer Validity | ... |
        | Type of Bid | Single/Two packet, evaluation method |
        | Evaluation Method | How bids are evaluated, price breakup requirements |
        | Estimated Bid Value (ECV) | Total amount with and without GST |
        | EMD (Earnest Money Deposit) | Amount, mode of payment, exemptions for MSE/Startup |
        | ePBG / PBG (Performance Bank Guarantee) | Percentage, duration, terms |
        | Payment Terms | Timeline, billing cycle, SDAC requirements |
        | Consignee / Reporting Officer | Name, address, designation |
        | MII / MSE Preference | Make in India and MSE purchase preference details |
        | Consortium / Sub-contracting | Whether allowed or not |
        | Arbitration / Mediation | Whether applicable |
        
        ### 02. SCOPE OF WORK
        Provide a detailed breakdown structured EXACTLY like this:
        (Write 1-2 paragraphs detailing exactly what is being procured, right-of-way details, and network lengths/jurisdictions)
        (Write a paragraph detailing the revenue heads or SOR components, e.g., SOR-A vs SOR-B caps, upkeep)
        
        | Work stream | What it obliges the contractor to do |
        |---|---|
        (List detailed work streams like "Preventive maintenance & patrolling", "Liaison & damage prevention", "Fault rectification", etc.)
        
        **Materials.** (Write a paragraph detailing exactly what the buyer supplies vs what is the contractor's responsibility)
        **Scope can move.** (Write a paragraph detailing scope variation clauses, length changes, etc.)
        
        **Penalty Schedule for SLA Non-Compliance:**
        (Provide a detailed table showing the exact penalties for missing SLAs, outages, cuts, etc.)
        | SLA Metric Missed | Penalty Imposed |
        |---|---|
        
        ### 03. BID STRATEGY & RECOMMENDATION
        Provide an expert analysis:
        a. **What the ECV Actually Buys (Monthly Breakdown)**
        Provide a detailed table allocating the monthly ECV by vertical (Service Head, Monthly Value, % of Monthly Revenue), followed by a Key Takeaway.
        b. **Price scenarios** — show a table of different discount levels and their impact.
        c. **Honest price assessment** — is this tender worth bidding on at the ECV considering minimum wages and actual costs?
        d. **Recommendation** — should {short_company} bid? At what price? What is the profit strategy?
        e. **MSE advantage** — how does our MSE/Startup status help here?
        
        ### 04. RISK ASSESSMENT
        Create a detailed risk matrix table:
        | Category | Risk Level | Assessment |
        |---|---|---|
        Cover these categories: Eligibility (technical), Eligibility (financial), Documentation, Execution/SLA, Cost escalation, Payment, Penalty exposure, Contractual, Timing.
        Risk levels: HIGH, MEDIUM, LOW, CRITICAL. Be honest and specific about each risk.
        
        ### 05. FINANCIAL ELIGIBILITY & TECHNICAL ELIGIBILITY
        Use HTML color tags for the status column to make it visually clear: 
        Use `<span style="color: green; font-weight: bold;">MATCH</span>` (we fully meet this), `<span style="color: #ffcc00; font-weight: bold;">PARTIAL</span>` (probably qualify but needs checking), or `<span style="color: red; font-weight: bold;">NO MATCH</span>` (we don't meet this).
        
        a. **FINANCIAL ELIGIBILITY** table:
        | What is Asked | How Much / What Exactly | What We Have | Status |
        |---|---|---|---|
        Cover: Average yearly turnover, value of past work, audited balance sheets, GST registration, company type, blacklisting status for {company_display}.
        
        b. **TECHNICAL ELIGIBILITY** table:
        | What is Asked | How Much / What Exactly | What We Have | Status |
        |---|---|---|---|
        Cover: Past experience requirements, manpower, equipment, deployment plans, certifications needed for {company_display}.
        
        ### 06. MANDATORY SUBMISSION CHECKLIST & DOCUMENT MAPPING
        Note on Mapping Logic: 
        - MAP: Exact document exists in the company library. 
        - CREATE: Document/declaration to be drafted on company letterhead / non-judicial stamp paper. 
        - MANUAL REQUIREMENT: Outside third-party issuance required (Bank, CA, Client).
        
        | # | Document as named in the ATC/tender | Action & File | Review & Operational Notes |
        |---|---|---|---|
        
        RULES for this section:
        - Copy document names WORD FOR WORD from the ATC/tender checklist provided
        - For Action & File: use 'MAP: [exact file path]' from {company_display}'s document library rules or 'CREATE: [explanation]' or 'MANUAL REQUIREMENT: [explanation]'.
        - The Review column must explain what to check before uploading, what to watch out for, and whether the bid will be rejected without this document
        - Include ALL documents from the checklist — do NOT skip any
        - If there are sub-items (like 2(i) and 2(ii)), each gets its own row
        
        CRITICAL REQUIREMENTS:
        1. Be EXHAUSTIVE. Extract every detail from the tender — do not summarize briefly.
        2. For Section 06, use the EXACT checklist from the ATC/tender document. Do NOT create your own list.
        3. Each checklist item MUST appear as a separate row.
        4. Map each document strictly against {company_display}'s Document Library Rules provided above to determine Action & file.
        5. The analysis should be detailed enough that someone who hasn't read the tender can understand everything.
        6. Include specific numbers, amounts, dates, names, and addresses wherever available.
        """
        
        contents = uploaded_files + [prompt]
        import time
        FALLBACK_MODELS = ["gemini-3.1-flash-lite", "gemini-3.5-flash-lite", "gemini-3.5-flash", "gemini-3.6-flash", "gemini-3.7-flash"]
        response = None
        for model_name in FALLBACK_MODELS:
            try:
                logging.info(f"Trying model: {model_name}")
                response = gemini_client.models.generate_content(model=model_name, contents=contents)
                logging.info(f"Success with model: {model_name}")
                break
            except Exception as e:
                logging.warning(f"Model {model_name} failed: {str(e)[:80]}")
                if model_name == FALLBACK_MODELS[-1]:
                    raise e
                time.sleep(2)
        
        html_content = response.text
        html_content = re.sub(r"^```(?:markdown|html)?\s*", "", html_content, flags=re.MULTILINE)
        html_content = re.sub(r"```\s*$", "", html_content, flags=re.MULTILINE)
        html_content = markdown.markdown(html_content, extensions=['tables'])
        
        template_path = Path("backend/core/pdf_template.html")
        if not template_path.exists():
            template_path = Path("core/pdf_template.html")
            
        if template_path.exists():
            template = template_path.read_text(encoding="utf-8")
        else:
            template = "<html><body>{{CONTENT}}</body></html>"
            
        tender_no = file.filename.replace(".pdf", "")
        full_html = template.replace("{{TENDER_NO}}", tender_no).replace("{{TENDER_TITLE}}", f"Tender Summary ({website.upper()} - {short_company})").replace("{{CONTENT}}", html_content)
        
        summary_name = f"Summary_{tender_no}_{website.upper()}_{short_company}.pdf"
        summary_path = config.DOWNLOADS_DIR / "summary" / summary_name
        summary_path.parent.mkdir(parents=True, exist_ok=True)
        await asyncio.to_thread(generate_pdf_from_html, full_html, str(summary_path))
        
        return {
            "status": "success",
            "message": f"Summary generated for {file.filename} ({website.upper()} - {short_company})",
            "summary_url": f"/downloads/summary/{summary_name}",
            "website": website,
            "company": company_display
        }
    except Exception as e:
        logging.error(f"Generate Summary error: {e}")
        return {"status": "error", "message": f"Failed to generate summary: {e}"}


from fastapi import Form

from typing import Optional, List
